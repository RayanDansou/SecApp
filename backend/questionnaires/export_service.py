"""
Service d'export des analyses de sécurité en PDF et DOCX.
Génère des rapports professionnels pour les questionnaires validés/rejetés.
"""

from io import BytesIO
from datetime import datetime
from django.template.loader import render_to_string
from weasyprint import HTML, CSS
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .models import QuestionnaireResponse, AIAnalysis


class ReportExportService:
    """Service pour générer des rapports d'analyse en PDF et DOCX."""

    def __init__(self, response_id):
        """
        Initialise le service avec une réponse au questionnaire.

        Args:
            response_id: ID de la QuestionnaireResponse à exporter
        """
        self.response = QuestionnaireResponse.objects.select_related(
            'questionnaire',
            'responder'
        ).prefetch_related(
            'answers__question',
            'ai_analyses',
            'comments__author',
            'status_history',
            'documents'
        ).get(id=response_id)

        # Récupérer la dernière analyse IA
        self.ai_analysis = self.response.ai_analyses.order_by('-created_at').first()

    def _get_context_data(self):
        """Prépare les données pour les templates."""
        return {
            'response': self.response,
            'questionnaire': self.response.questionnaire,
            'responder': self.response.responder,
            'ai_analysis': self.ai_analysis,
            'answers': self.response.answers.select_related('question').order_by('question__order'),
            'comments': self.response.comments.select_related('author').order_by('created_at'),
            'documents': self.response.documents.all(),
            'status_history': self.response.status_history.select_related('changed_by').order_by('changed_at'),
            'generated_at': datetime.now(),
            'status_badge_color': self._get_status_color(),
        }

    def _get_status_color(self):
        """Retourne la couleur du badge selon le statut."""
        colors = {
            'VALIDE': '#10b981',  # Green
            'REJETE': '#ef4444',  # Red
            'EN_VALIDATION': '#f59e0b',  # Orange
            'SOUMIS': '#3b82f6',  # Blue
            'BROUILLON': '#6b7280',  # Gray
        }
        return colors.get(self.response.status, '#6b7280')

    def generate_pdf(self):
        """
        Génère un rapport PDF.

        Returns:
            BytesIO: Contenu du PDF
        """
        context = self._get_context_data()

        # Rendre le template HTML
        html_string = render_to_string('questionnaires/export/report_template.html', context)

        # Générer le PDF depuis HTML
        html = HTML(string=html_string)
        pdf_file = BytesIO()
        html.write_pdf(pdf_file, stylesheets=[CSS(string=self._get_pdf_styles())])
        pdf_file.seek(0)

        return pdf_file

    def _get_pdf_styles(self):
        """Retourne les styles CSS pour le PDF."""
        return """
            @page {
                size: A4;
                margin: 2cm;
                @top-center {
                    content: "Rapport d'Analyse de Sécurité - GuardianIQ";
                    font-size: 10px;
                    color: #666;
                }
                @bottom-right {
                    content: "Page " counter(page) " sur " counter(pages);
                    font-size: 10px;
                    color: #666;
                }
            }

            body {
                font-family: 'Arial', sans-serif;
                line-height: 1.6;
                color: #333;
            }

            .header {
                text-align: center;
                margin-bottom: 30px;
                padding-bottom: 20px;
                border-bottom: 3px solid #3b82f6;
            }

            .header h1 {
                color: #1e40af;
                font-size: 28px;
                margin-bottom: 10px;
            }

            .header .subtitle {
                color: #6b7280;
                font-size: 14px;
            }

            .status-badge {
                display: inline-block;
                padding: 8px 16px;
                border-radius: 20px;
                color: white;
                font-weight: bold;
                font-size: 14px;
                margin: 10px 0;
            }

            .section {
                margin-bottom: 30px;
                page-break-inside: avoid;
            }

            .section-title {
                color: #1e40af;
                font-size: 20px;
                border-bottom: 2px solid #e5e7eb;
                padding-bottom: 10px;
                margin-bottom: 15px;
            }

            .info-grid {
                display: grid;
                grid-template-columns: 1fr 1fr;
                gap: 15px;
                margin-bottom: 20px;
            }

            .info-item {
                padding: 10px;
                background: #f9fafb;
                border-left: 3px solid #3b82f6;
            }

            .info-label {
                font-weight: bold;
                color: #4b5563;
                font-size: 12px;
                text-transform: uppercase;
            }

            .info-value {
                color: #111827;
                font-size: 14px;
                margin-top: 5px;
            }

            .score-container {
                display: grid;
                grid-template-columns: 1fr 1fr 1fr 1fr;
                gap: 15px;
                margin: 20px 0;
            }

            .score-card {
                text-align: center;
                padding: 20px;
                background: #f9fafb;
                border-radius: 8px;
                border: 2px solid #e5e7eb;
            }

            .score-value {
                font-size: 36px;
                font-weight: bold;
                color: #1e40af;
            }

            .score-label {
                font-size: 12px;
                color: #6b7280;
                text-transform: uppercase;
                margin-top: 5px;
            }

            .question-item {
                margin-bottom: 20px;
                padding: 15px;
                background: #f9fafb;
                border-radius: 8px;
                page-break-inside: avoid;
            }

            .question-text {
                font-weight: bold;
                color: #1e40af;
                margin-bottom: 8px;
            }

            .answer-text {
                color: #374151;
                padding-left: 15px;
                border-left: 3px solid #3b82f6;
            }

            .recommendation-item {
                padding: 10px 15px;
                margin-bottom: 10px;
                background: #fef3c7;
                border-left: 4px solid #f59e0b;
                border-radius: 4px;
            }

            .strength-item {
                padding: 10px 15px;
                margin-bottom: 10px;
                background: #d1fae5;
                border-left: 4px solid #10b981;
                border-radius: 4px;
            }

            .weakness-item {
                padding: 10px 15px;
                margin-bottom: 10px;
                background: #fee2e2;
                border-left: 4px solid #ef4444;
                border-radius: 4px;
            }

            .comment-item {
                margin-bottom: 15px;
                padding: 12px;
                background: #f3f4f6;
                border-radius: 6px;
                page-break-inside: avoid;
            }

            .comment-author {
                font-weight: bold;
                color: #1e40af;
            }

            .comment-date {
                font-size: 12px;
                color: #6b7280;
            }

            .footer {
                margin-top: 40px;
                padding-top: 20px;
                border-top: 2px solid #e5e7eb;
                text-align: center;
                font-size: 12px;
                color: #6b7280;
            }

            table {
                width: 100%;
                border-collapse: collapse;
                margin: 15px 0;
            }

            th, td {
                padding: 12px;
                text-align: left;
                border-bottom: 1px solid #e5e7eb;
            }

            th {
                background: #f3f4f6;
                font-weight: bold;
                color: #1e40af;
            }
        """

    def generate_docx(self):
        """
        Génère un rapport Word (DOCX).

        Returns:
            BytesIO: Contenu du document DOCX
        """
        doc = Document()
        context = self._get_context_data()

        # Configuration du document
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # En-tête
        self._add_docx_header(doc, context)

        # Informations générales
        self._add_docx_general_info(doc, context)

        # Scores CIA
        if self.ai_analysis:
            self._add_docx_cia_scores(doc, context)

        # Questions et Réponses
        self._add_docx_qa_section(doc, context)

        # Analyse IA
        if self.ai_analysis:
            self._add_docx_ai_analysis(doc, context)

        # Commentaires
        if context['comments']:
            self._add_docx_comments(doc, context)

        # Historique
        self._add_docx_history(doc, context)

        # Pied de page
        self._add_docx_footer(doc, context)

        # Sauvegarder dans BytesIO
        docx_file = BytesIO()
        doc.save(docx_file)
        docx_file.seek(0)

        return docx_file

    def _add_docx_header(self, doc, context):
        """Ajoute l'en-tête du document DOCX."""
        # Titre principal
        title = doc.add_heading('Rapport d\'Analyse de Sécurité', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].font.color.rgb = RGBColor(30, 64, 175)

        # Sous-titre
        subtitle = doc.add_paragraph('GuardianIQ - Analyse de Sécurité des Applications')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.size = Pt(12)
        subtitle.runs[0].font.color.rgb = RGBColor(107, 114, 128)

        # Statut
        status_para = doc.add_paragraph()
        status_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        status_run = status_para.add_run(f'Statut: {context["response"].get_status_display()}')
        status_run.bold = True
        status_run.font.size = Pt(14)

        doc.add_paragraph()  # Espacement

    def _add_docx_general_info(self, doc, context):
        """Ajoute les informations générales."""
        doc.add_heading('Informations Générales', 1)

        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'

        rows_data = [
            ('Questionnaire', context['questionnaire'].title),
            ('Chef de Projet', context['responder'].get_full_name() or context['responder'].username),
            ('Date de soumission', context['response'].submitted_at.strftime('%d/%m/%Y %H:%M') if context['response'].submitted_at else 'N/A'),
            ('Statut', context['response'].get_status_display()),
            ('Date de génération', context['generated_at'].strftime('%d/%m/%Y %H:%M')),
        ]

        for i, (label, value) in enumerate(rows_data):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[0].paragraphs[0].runs[0].bold = True
            table.rows[i].cells[1].text = str(value)

        doc.add_paragraph()

    def _add_docx_cia_scores(self, doc, context):
        """Ajoute les scores CIA."""
        doc.add_heading('Scores de Sécurité (CIA)', 1)

        ai = context['ai_analysis']

        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'

        scores = [
            ('Score de Cohérence Global', f'{ai.coherence_score}%'),
            ('Confidentialité (C)', f'{ai.confidentiality_score}/100'),
            ('Intégrité (I)', f'{ai.integrity_score}/100'),
            ('Disponibilité (A)', f'{ai.availability_score}/100'),
        ]

        for i, (label, score) in enumerate(scores):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[0].paragraphs[0].runs[0].bold = True
            score_run = table.rows[i].cells[1].paragraphs[0].add_run(score)
            score_run.bold = True
            score_run.font.size = Pt(14)
            score_run.font.color.rgb = RGBColor(30, 64, 175)

        doc.add_paragraph()

    def _add_docx_qa_section(self, doc, context):
        """Ajoute la section Questions & Réponses."""
        doc.add_heading('Questions et Réponses', 1)

        for answer in context['answers']:
            # Question
            q_para = doc.add_paragraph()
            q_run = q_para.add_run(f'Q{answer.question.order}: {answer.question.text}')
            q_run.bold = True
            q_run.font.color.rgb = RGBColor(30, 64, 175)

            # Réponse
            a_para = doc.add_paragraph(answer.answer_text, style='List Bullet')
            a_para.paragraph_format.left_indent = Inches(0.5)

            doc.add_paragraph()  # Espacement

    def _add_docx_ai_analysis(self, doc, context):
        """Ajoute l'analyse IA."""
        ai = context['ai_analysis']

        # Résumé
        doc.add_heading('Résumé de l\'Analyse IA', 1)
        doc.add_paragraph(ai.analysis_summary)
        doc.add_paragraph()

        # Points forts
        if ai.strengths:
            doc.add_heading('Points Forts', 2)
            for strength in ai.strengths:
                doc.add_paragraph(strength, style='List Bullet')
            doc.add_paragraph()

        # Points faibles
        if ai.weaknesses:
            doc.add_heading('Points Faibles', 2)
            for weakness in ai.weaknesses:
                doc.add_paragraph(weakness, style='List Bullet')
            doc.add_paragraph()

        # Recommandations
        if ai.recommendations:
            doc.add_heading('Recommandations', 2)
            for rec in ai.recommendations:
                doc.add_paragraph(rec, style='List Bullet')
            doc.add_paragraph()

        # Incohérences
        if ai.inconsistencies:
            doc.add_heading('Incohérences Détectées', 2)
            for inc in ai.inconsistencies:
                doc.add_paragraph(inc, style='List Bullet')
            doc.add_paragraph()

    def _add_docx_comments(self, doc, context):
        """Ajoute les commentaires."""
        doc.add_heading('Commentaires', 1)

        for comment in context['comments']:
            p = doc.add_paragraph()
            p.add_run(f'{comment.author.get_full_name() or comment.author.username}: ').bold = True
            p.add_run(f'{comment.content}')
            p.add_run(f'\n({comment.created_at.strftime("%d/%m/%Y %H:%M")})').font.size = Pt(9)
            doc.add_paragraph()

    def _add_docx_history(self, doc, context):
        """Ajoute l'historique des statuts."""
        doc.add_heading('Historique des Statuts', 1)

        table = doc.add_table(rows=len(context['status_history']) + 1, cols=4)
        table.style = 'Light Grid Accent 1'

        # En-têtes
        headers = ['Date', 'Ancien Statut', 'Nouveau Statut', 'Modifié par']
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True

        # Données
        for i, history in enumerate(context['status_history'], 1):
            table.rows[i].cells[0].text = history.changed_at.strftime('%d/%m/%Y %H:%M')
            table.rows[i].cells[1].text = history.get_old_status_display() or 'Création'
            table.rows[i].cells[2].text = history.get_new_status_display()
            table.rows[i].cells[3].text = history.changed_by.username if history.changed_by else 'Système'

        doc.add_paragraph()

    def _add_docx_footer(self, doc, context):
        """Ajoute le pied de page."""
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run(
            f'Rapport généré par GuardianIQ le {context["generated_at"].strftime("%d/%m/%Y à %H:%M")}\n'
            'Ce document est confidentiel et destiné uniquement aux personnes autorisées.'
        )
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor(107, 114, 128)

    def get_filename(self, extension):
        """
        Génère un nom de fichier pour l'export.

        Args:
            extension: 'pdf' ou 'docx'

        Returns:
            str: Nom du fichier
        """
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        questionnaire_title = self.response.questionnaire.title.replace(' ', '_')[:30]
        return f'Rapport_{questionnaire_title}_{timestamp}.{extension}'
